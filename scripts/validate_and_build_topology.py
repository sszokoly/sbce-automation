#!/usr/bin/env python3

from __future__ import annotations

import argparse
import ipaddress
import json
import re
import socket
import sys
from pathlib import Path
from typing import Any, Iterable, Literal, Optional
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import openpyxl
from pydantic import BaseModel, ConfigDict, Field, ValidationError, field_validator, model_validator


DOMAIN_LABEL_RE = re.compile(r"^[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?$")
PASSWORD_SPECIALS = set("-_@*!")
PASSWORD_FORBIDDEN = set("#$&")


class ConfigError(Exception):
    pass


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def split_list(value: str) -> list[str]:
    return [item.strip() for item in value.split(",") if item.strip()]


def parse_target_sheets(value: Optional[str]) -> Optional[set[str]]:
    if not value:
        return None
    result = {item.strip() for item in value.split(",") if item.strip()}
    return result or None


def fail(message: str) -> None:
    raise ValueError(message)


def require_string(name: str, value: str, max_len: Optional[int] = None) -> str:
    if not value:
        fail(f"'{name}' cannot be empty")
    if max_len is not None and len(value) > max_len:
        fail(f"'{name}' must be up to {max_len} characters")
    return value


def require_enum(name: str, value: str, choices: Iterable[str]) -> str:
    normalized = require_string(name, value).upper()
    allowed = set(choices)
    if normalized not in allowed:
        fail(f"'{name}' must be one of: {', '.join(sorted(allowed))}")
    return normalized


def require_lower_enum(name: str, value: str, choices: Iterable[str]) -> str:
    normalized = require_string(name, value).lower()
    allowed = set(choices)
    if normalized not in allowed:
        fail(f"'{name}' must be one of: {', '.join(sorted(allowed))}")
    return normalized


def validate_ipv4(name: str, value: str) -> None:
    try:
        ipaddress.IPv4Address(value)
    except ValueError:
        fail(f"'{name}' must be a valid IPv4 address")


def validate_ipv6(name: str, value: str) -> None:
    try:
        ipaddress.IPv6Address(value)
    except ValueError:
        fail(f"'{name}' must be a valid IPv6 address")


def validate_ip(name: str, value: str) -> None:
    try:
        ipaddress.ip_address(value)
    except ValueError:
        fail(f"'{name}' must be a valid IPv4 or IPv6 address")


def validate_ipv4_subnet_mask(name: str, value: str) -> None:
    if value.startswith("/"):
        try:
            prefix = int(value[1:])
        except ValueError:
            fail(f"'{name}' must be a valid IPv4 subnet mask")
        if prefix < 0 or prefix > 32:
            fail(f"'{name}' must be an IPv4 prefix between /0 and /32")
        return

    try:
        mask_int = int(ipaddress.IPv4Address(value))
    except ValueError:
        fail(f"'{name}' must be a valid IPv4 subnet mask")

    inverse = (~mask_int) & 0xFFFFFFFF
    if inverse & (inverse + 1) != 0:
        fail(f"'{name}' must be a contiguous IPv4 subnet mask")


def validate_ipv6_prefix(name: str, value: str) -> None:
    raw = value[1:] if value.startswith("/") else value
    try:
        prefix = int(raw)
    except ValueError:
        fail(f"'{name}' must be an IPv6 prefix between 0 and 128")
    if prefix < 0 or prefix > 128:
        fail(f"'{name}' must be an IPv6 prefix between 0 and 128")


def validate_sig_mask(name: str, value: str) -> None:
    try:
        validate_ipv4_subnet_mask(name, value)
        return
    except ValueError:
        pass
    validate_ipv6_prefix(name, value)


def is_domain_suffix(value: str) -> bool:
    if not value or value.startswith(".") or value.endswith(".") or "." not in value or len(value) > 253:
        return False
    return all(DOMAIN_LABEL_RE.match(label) for label in value.split("."))


def validate_domain_suffix(name: str, value: str) -> None:
    if not is_domain_suffix(value):
        fail(f"'{name}' must be a valid domain suffix")


def validate_ipv4_or_fqdn(name: str, value: str) -> None:
    try:
        ipaddress.IPv4Address(value)
        return
    except ValueError:
        pass
    if not is_domain_suffix(value):
        fail(f"'{name}' must be a valid IPv4 address or FQDN")


def validate_resolvable_ipv4_or_fqdn(name: str, value: str) -> None:
    try:
        ipaddress.IPv4Address(value)
        return
    except ValueError:
        pass
    if not is_domain_suffix(value):
        fail(f"'{name}' must be a valid IPv4 address or resolvable FQDN")
    try:
        socket.getaddrinfo(value, None, family=socket.AF_INET)
    except socket.gaierror:
        fail(f"'{name}' FQDN must resolve to an IPv4 address from the controller host")


def validate_timezone(name: str, value: str) -> None:
    try:
        ZoneInfo(value)
    except ZoneInfoNotFoundError:
        fail(f"'{name}' must be a valid IANA time zone")


def validate_password(name: str, value: str) -> None:
    require_string(name, value)
    if len(value) < 8:
        fail(f"'{name}' must be at least 8 characters")
    if not any(ch.isupper() for ch in value):
        fail(f"'{name}' must contain at least 1 uppercase letter")
    if not any(ch.islower() for ch in value):
        fail(f"'{name}' must contain at least 1 lowercase letter")
    if not any(ch.isdigit() for ch in value):
        fail(f"'{name}' must contain at least 1 number")
    if not any(ch in PASSWORD_SPECIALS for ch in value):
        fail("'%s' must contain at least 1 special character from: - _ @ * !" % name)
    forbidden = sorted({ch for ch in value if ch in PASSWORD_FORBIDDEN})
    if forbidden:
        fail(f"'{name}' cannot contain: {' '.join(forbidden)}")


class SbceSheet(BaseModel):
    model_config = ConfigDict(extra="allow", populate_by_name=True)

    sheet_name: str = Field(alias="_sheet")
    sheet_index: int = Field(alias="_index")
    platform_type: str = ""
    platform_address: str = ""
    platform_username: str = ""
    platform_password: str = ""
    datastore: str = ""
    ovf: str = ""
    vmname: str = ""
    ipmode: str = ""
    hostname: str = ""
    apptype: str = ""
    appname: str = ""
    nwpass: str = ""
    ems_inst_type: str = ""
    ip0: str = ""
    netmask0: str = ""
    gateway: str = ""
    ipv6address0: str = ""
    ipv6prefix0: str = ""
    ipv6gateway: str = ""
    domain_suffix: str = ""
    first_last_name: str = ""
    organizational_unit: str = ""
    organization: str = ""
    locality: str = ""
    state_province: str = ""
    country: str = ""
    timezone: str = ""
    ntpservers: str = ""
    ntpipv6: str = ""
    dns: str = ""
    emsip: str = ""
    emsip_v6: str = ""
    rootpass: str = ""
    ipcspass: str = ""
    grubpass: str = ""
    ucsecpass: str = ""
    M1: str = ""
    M2: str = ""
    A1: str = ""
    A2: str = ""
    B1: str = ""
    B2: str = ""
    ha_peer_node: str = ""
    ha_peer_ip: str = ""
    sig_iface: str = ""
    sig_name: str = ""
    sig_mask: str = ""
    sig_gw: str = ""
    sig_ip: str = ""

    @field_validator("*", mode="before")
    @classmethod
    def normalize_value(cls, value: Any) -> Any:
        if isinstance(value, int):
            return value
        return text(value)

    @model_validator(mode="after")
    def validate_fields(self) -> "SbceSheet":
        self.platform_type = require_enum("platform_type", self.platform_type, ("KVM", "ESXI"))
        validate_resolvable_ipv4_or_fqdn("platform_address", require_string("platform_address", self.platform_address))
        require_string("platform_username", self.platform_username)
        require_string("platform_password", self.platform_password)
        require_string("datastore", self.datastore)
        require_string("ovf", self.ovf)
        require_string("vmname", self.vmname)
        self.ipmode = require_enum("ipmode", self.ipmode, ("DUAL_STACK", "IPV4"))
        require_string("hostname", self.hostname, max_len=20)
        self.apptype = require_enum("apptype", self.apptype, ("EMS+SBCE", "EMS", "SBCE"))
        require_string("appname", self.appname, max_len=20)
        if self.apptype == "SBCE" and self.appname == self.vmname:
            fail("'appname' cannot be the same as 'vmname' when 'apptype' is 'SBCE'")
        if not self.nwpass:
            self.nwpass = "avaya"

        if self.apptype == "EMS":
            self.ems_inst_type = require_lower_enum("ems_inst_type", self.ems_inst_type, ("primary", "secondary"))
        else:
            self.ems_inst_type = ""

        if not self.ip0 and not self.ipv6address0:
            fail("'ip0' cannot be empty when 'ipv6address0' is empty")
        if self.ip0:
            validate_ipv4("ip0", self.ip0)
        if not self.netmask0 and not self.ipv6prefix0:
            fail("'netmask0' cannot be empty when 'ipv6prefix0' is empty")
        if self.netmask0:
            validate_ipv4_subnet_mask("netmask0", self.netmask0)
        if not self.gateway and not self.ipv6gateway:
            fail("'gateway' cannot be empty when 'ipv6gateway' is empty")
        if self.gateway:
            validate_ipv4("gateway", self.gateway)
        if self.ipv6address0:
            validate_ipv6("ipv6address0", self.ipv6address0)
        if self.ipv6prefix0:
            validate_ipv6_prefix("ipv6prefix0", self.ipv6prefix0)
        if self.ipv6gateway:
            validate_ipv6("ipv6gateway", self.ipv6gateway)

        validate_domain_suffix("domain_suffix", require_string("domain_suffix", self.domain_suffix))
        expected_fln = f"{self.hostname}.{self.domain_suffix}"
        if require_string("first_last_name", self.first_last_name) != expected_fln:
            fail(f"'first_last_name' must be '{expected_fln}'")

        if self.apptype != "SBCE":
            require_string("organizational_unit", self.organizational_unit)
            require_string("organization", self.organization)
            require_string("locality", self.locality)
            require_string("state_province", self.state_province)
            require_string("country", self.country)

        validate_timezone("timezone", require_string("timezone", self.timezone))

        if self.ntpservers:
            items = split_list(self.ntpservers)
            if not items:
                fail("'ntpservers' must be a valid IPv4 address or FQDN")
            for item in items:
                validate_ipv4_or_fqdn("ntpservers", item)
        if self.ntpipv6:
            validate_ipv6("ntpipv6", self.ntpipv6)
        validate_ipv4("dns", require_string("dns", self.dns))

        if self.emsip:
            validate_ipv4("emsip", self.emsip)
        if self.emsip_v6:
            validate_ipv6("emsip_v6", self.emsip_v6)
        needs_ems_ip = self.apptype == "SBCE" or (self.apptype == "EMS" and self.ems_inst_type == "secondary")
        if needs_ems_ip and not self.emsip and not self.emsip_v6:
            fail("'emsip' or 'emsip_v6' cannot be empty for SBCE or secondary EMS sheets")

        validate_password("rootpass", self.rootpass)
        validate_password("ipcspass", self.ipcspass)
        validate_password("grubpass", self.grubpass)
        validate_password("ucsecpass", self.ucsecpass)

        require_string("M1", self.M1)
        if self.apptype == "EMS" or (self.apptype == "SBCE" and self.ha_peer_node):
            require_string("M2", self.M2)

        if self.apptype != "EMS":
            require_string("A1", self.A1)
            require_string("A2", self.A2)
            require_string("B1", self.B1)
            require_string("B2", self.B2)

        if self.ha_peer_ip and not self.ha_peer_node:
            fail("'ha_peer_node' cannot be empty when 'ha_peer_ip' is set")
        if self.ha_peer_node and not self.ha_peer_ip:
            fail("'ha_peer_ip' cannot be empty when 'ha_peer_node' is set")
        if self.ha_peer_ip:
            validate_ip("ha_peer_ip", self.ha_peer_ip)

        if self.sig_iface or self.apptype != "EMS":
            self.sig_iface = require_enum("sig_iface", self.sig_iface, ("A1", "A2", "B1", "B2"))
        if self.sig_name or self.apptype != "EMS":
            require_string("sig_name", self.sig_name, max_len=20)
        if self.sig_mask or self.apptype != "EMS":
            validate_sig_mask("sig_mask", require_string("sig_mask", self.sig_mask))
        if self.sig_gw or self.apptype != "EMS":
            validate_ip("sig_gw", require_string("sig_gw", self.sig_gw))
        if self.sig_ip or self.apptype != "EMS":
            validate_ip("sig_ip", require_string("sig_ip", self.sig_ip))
            if self.ip0 and self.sig_ip == self.ip0:
                fail("'sig_ip' cannot be the same as 'ip0'")

        return self

    @property
    def is_ha(self) -> bool:
        return bool(self.ha_peer_node and self.ha_peer_ip)

    @property
    def management_ip(self) -> str:
        return self.ip0 or self.ipv6address0

    def topology_dict(self, artifact_role: Optional[str] = None) -> dict[str, Any]:
        data = self.model_dump(by_alias=True)
        if artifact_role is not None:
            data["artifact_role"] = artifact_role
        return data


def first_validation_error(exc: ValidationError) -> str:
    error = exc.errors()[0]
    ctx_error = error.get("ctx", {}).get("error")
    if ctx_error is not None:
        return str(ctx_error)
    location = ".".join(str(part) for part in error.get("loc", []))
    message = error.get("msg", str(exc))
    return f"{location}: {message}" if location else message


def load_workbook(path: Path) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheets: list[dict[str, Any]] = []
    for index, ws in enumerate(workbook.worksheets):
        data: dict[str, Any] = {}
        for row in ws.iter_rows(values_only=True):
            key = text(row[0] if row else None)
            if key:
                data[key] = text(row[1] if len(row) > 1 else None)
        data["_sheet"] = ws.title
        data["_index"] = index
        sheets.append(data)
    return sheets


def validate_sheets(raw_sheets: list[dict[str, Any]], selected_names: Optional[set[str]]) -> list[SbceSheet]:
    all_names = {sheet["_sheet"] for sheet in raw_sheets}
    if selected_names:
        for name in sorted(selected_names - all_names):
            raise ConfigError(f"target_sheets: unknown sheet '{name}'")
        raw_sheets = [sheet for sheet in raw_sheets if sheet["_sheet"] in selected_names]

    sheets: list[SbceSheet] = []
    for raw_sheet in raw_sheets:
        try:
            sheets.append(SbceSheet.model_validate(raw_sheet))
        except ValidationError as exc:
            raise ConfigError(f"{raw_sheet['_sheet']}: {first_validation_error(exc)}") from exc
    return sheets


def ip_sort_key(value: Any) -> tuple[int, int | str]:
    raw = text(value)
    try:
        return (0, int(ipaddress.ip_address(raw)))
    except ValueError:
        return (1, raw)


def sheet_sort_key(sheet: SbceSheet) -> tuple[int, int | str]:
    return ip_sort_key(sheet.management_ip)


def ensure_unique(sheets: list[SbceSheet], field: Literal["ip0", "hostname", "vmname"]) -> None:
    seen: dict[str, SbceSheet] = {}
    for sheet in sheets:
        value = getattr(sheet, field)
        if not value:
            continue
        previous = seen.get(value)
        if previous:
            raise ConfigError(f"{sheet.sheet_name}: duplicate {field} '{value}' also used by {previous.sheet_name}")
        seen[value] = sheet


def build_management_ip_index(sheets: list[SbceSheet]) -> dict[str, SbceSheet]:
    by_ip: dict[str, SbceSheet] = {}
    for sheet in sheets:
        for value in (sheet.ip0, sheet.ipv6address0):
            if not value:
                continue
            previous = by_ip.get(value)
            if previous:
                raise ConfigError(f"{sheet.sheet_name}: management address '{value}' also used by {previous.sheet_name}")
            by_ip[value] = sheet
    return by_ip


def resolve_peer_node(sheet: SbceSheet, sheets: list[SbceSheet]) -> Optional[SbceSheet]:
    for candidate in sheets:
        if sheet.ha_peer_node in {candidate.hostname, candidate.sheet_name, candidate.vmname}:
            return candidate
    return None


def validate_ha_pairs(sheets: list[SbceSheet], by_management_ip: dict[str, SbceSheet]) -> dict[str, SbceSheet]:
    peers: dict[str, SbceSheet] = {}
    for sheet in sheets:
        if not sheet.is_ha:
            continue
        if sheet.apptype != "SBCE":
            raise ConfigError(f"{sheet.sheet_name}: only SBCE sheets can be paired for HA")

        peer_by_node = resolve_peer_node(sheet, sheets)
        peer_by_ip = by_management_ip.get(sheet.ha_peer_ip)
        if not peer_by_node:
            raise ConfigError(f"{sheet.sheet_name}: ha_peer_node '{sheet.ha_peer_node}' does not match a selected sheet")
        if not peer_by_ip:
            raise ConfigError(f"{sheet.sheet_name}: ha_peer_ip '{sheet.ha_peer_ip}' does not match a selected sheet management address")
        if peer_by_node is not peer_by_ip:
            raise ConfigError(f"{sheet.sheet_name}: ha_peer_node and ha_peer_ip identify different sheets")
        if peer_by_node.apptype != "SBCE":
            raise ConfigError(f"{sheet.sheet_name}: HA peer '{peer_by_node.sheet_name}' is not an SBCE sheet")
        if not peer_by_node.is_ha:
            raise ConfigError(f"{sheet.sheet_name}: HA peer '{peer_by_node.sheet_name}' does not point back to this node")

        peer_target_by_node = resolve_peer_node(peer_by_node, sheets)
        peer_target_by_ip = by_management_ip.get(peer_by_node.ha_peer_ip)
        if peer_target_by_node is not sheet or peer_target_by_ip is not sheet:
            raise ConfigError(f"{sheet.sheet_name}: HA peer '{peer_by_node.sheet_name}' does not point back to this node")
        if sheet.appname != peer_by_node.appname:
            raise ConfigError(f"{sheet.sheet_name}: HA peer '{peer_by_node.sheet_name}' must use the same appname")
        if sheet.sig_ip != peer_by_node.sig_ip:
            raise ConfigError(f"{sheet.sheet_name}: HA peer '{peer_by_node.sheet_name}' must use the same sig_ip")
        peers[sheet.sheet_name] = peer_by_node
    return peers


def are_ha_peers(left: SbceSheet, right: SbceSheet, peers: dict[str, SbceSheet]) -> bool:
    return peers.get(left.sheet_name) is right and peers.get(right.sheet_name) is left


def ensure_unique_or_ha_pair(sheets: list[SbceSheet], field: Literal["appname", "sig_ip"], peers: dict[str, SbceSheet]) -> None:
    seen: dict[str, SbceSheet] = {}
    for sheet in sheets:
        value = getattr(sheet, field)
        if not value:
            continue
        previous = seen.get(value)
        if previous and not are_ha_peers(previous, sheet, peers):
            raise ConfigError(f"{sheet.sheet_name}: duplicate {field} '{value}' also used by {previous.sheet_name}")
        seen[value] = sheet


def ensure_sig_ip_not_management_ip(sheets: list[SbceSheet], by_management_ip: dict[str, SbceSheet]) -> None:
    for sheet in sheets:
        if not sheet.sig_ip:
            continue
        owner = by_management_ip.get(sheet.sig_ip)
        if owner:
            raise ConfigError(f"{sheet.sheet_name}: sig_ip '{sheet.sig_ip}' conflicts with management address on {owner.sheet_name}")


def primary_for_child(sheet: SbceSheet, primary_by_ip: dict[str, SbceSheet]) -> SbceSheet:
    matches: list[tuple[str, SbceSheet]] = []
    for field, value in (("emsip", sheet.emsip), ("emsip_v6", sheet.emsip_v6)):
        if not value:
            continue
        primary = primary_by_ip.get(value)
        if not primary:
            raise ConfigError(f"{sheet.sheet_name}: {field} '{value}' does not match a selected primary EMS management address")
        matches.append((field, primary))

    if not matches:
        raise ConfigError(f"{sheet.sheet_name}: emsip or emsip_v6 must match a selected primary EMS")

    primary = matches[0][1]
    for _, candidate in matches[1:]:
        if candidate is not primary:
            raise ConfigError(f"{sheet.sheet_name}: emsip and emsip_v6 refer to different primary EMS sheets")
    return primary


def validate_cross_sheet_restrictions(sheets: list[SbceSheet]) -> tuple[str, dict[str, SbceSheet]]:
    platforms = {sheet.platform_type for sheet in sheets}
    if len(platforms) != 1:
        raise ConfigError("selected sheets must contain exactly one platform_type")

    ensure_unique(sheets, "ip0")
    ensure_unique(sheets, "hostname")
    ensure_unique(sheets, "vmname")
    by_management_ip = build_management_ip_index(sheets)
    peers = validate_ha_pairs(sheets, by_management_ip)
    ensure_unique_or_ha_pair(sheets, "appname", peers)
    ensure_unique_or_ha_pair(sheets, "sig_ip", peers)
    ensure_sig_ip_not_management_ip(sheets, by_management_ip)
    return platforms.pop(), peers


def build_topology(sheets: list[SbceSheet], platform: str, peers: dict[str, SbceSheet]) -> dict[str, Any]:
    primary_groups: dict[str, dict[str, Any]] = {}
    primary_by_ip: dict[str, SbceSheet] = {}

    for sheet in sheets:
        if sheet.apptype == "EMS+SBCE" or (sheet.apptype == "EMS" and sheet.ems_inst_type != "secondary"):
            primary_groups[sheet.sheet_name] = {
                "primary_sheet": sheet.sheet_name,
                "primary": sheet.topology_dict(artifact_role="EMS1"),
                "coresident": sheet.apptype == "EMS+SBCE",
                "secondary_ems": [],
                "sbce_units": [],
            }
            for value in (sheet.ip0, sheet.ipv6address0):
                if value:
                    primary_by_ip[value] = sheet

    if not primary_groups:
        raise ConfigError("at least one primary EMS or EMS+SBCE sheet is required")

    child_primary_by_sheet: dict[str, SbceSheet] = {}
    for sheet in sheets:
        if sheet.apptype == "SBCE" or (sheet.apptype == "EMS" and sheet.ems_inst_type == "secondary"):
            primary = primary_for_child(sheet, primary_by_ip)
            child_primary_by_sheet[sheet.sheet_name] = primary
            group = primary_groups[primary.sheet_name]
            if sheet.apptype == "EMS":
                group["secondary_ems"].append(sheet.topology_dict(artifact_role="EMS2"))

    for primary in [group["primary"] for group in primary_groups.values()]:
        primary_sheet = primary["_sheet"]
        group = primary_groups[primary_sheet]
        sbces = [
            sheet for sheet in sheets
            if sheet.apptype == "SBCE" and child_primary_by_sheet.get(sheet.sheet_name) is not None
            and child_primary_by_sheet[sheet.sheet_name].sheet_name == primary_sheet
        ]
        sbces.sort(key=sheet_sort_key)
        sbce_roles = {sheet.sheet_name: f"SBC{index}" for index, sheet in enumerate(sbces, start=1)}

        consumed: set[str] = set()
        for sheet in sbces:
            if sheet.sheet_name in consumed:
                continue
            peer = peers.get(sheet.sheet_name)
            if peer:
                if child_primary_by_sheet.get(peer.sheet_name) is not child_primary_by_sheet[sheet.sheet_name]:
                    raise ConfigError(f"{sheet.sheet_name}: HA peer '{peer.sheet_name}' belongs to a different EMS")
                ordered = sorted([sheet, peer], key=sheet_sort_key)
                group["sbce_units"].append({
                    "type": "ha",
                    "sheets": [node.sheet_name for node in ordered],
                    "nodes": [node.topology_dict(artifact_role=sbce_roles[node.sheet_name]) for node in ordered],
                })
                consumed.update({sheet.sheet_name, peer.sheet_name})
            else:
                group["sbce_units"].append({
                    "type": "single",
                    "sheets": [sheet.sheet_name],
                    "nodes": [sheet.topology_dict(artifact_role=sbce_roles[sheet.sheet_name])],
                })
                consumed.add(sheet.sheet_name)

    groups = list(primary_groups.values())
    groups.sort(key=lambda group: ip_sort_key(group["primary"].get("ip0") or group["primary"].get("ipv6address0")))
    return {"ok": True, "platform": platform, "groups": groups}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Validate SBCE workbook and build deployment topology")
    parser.add_argument("src", help="Workbook path")
    parser.add_argument("--target-sheets", default="", help="Comma-separated sheet names to execute exactly")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        raw_sheets = load_workbook(Path(args.src))
    except Exception as exc:
        print(json.dumps({"ok": False, "errors": [str(exc)]}, indent=2))
        return 1

    try:
        selected_names = parse_target_sheets(args.target_sheets)
        sheets = validate_sheets(raw_sheets, selected_names)
        platform, peers = validate_cross_sheet_restrictions(sheets)
        topology = build_topology(sheets, platform, peers)
    except (ConfigError, ValueError) as exc:
        print(json.dumps({"ok": False, "errors": [str(exc)]}, indent=2))
        return 1

    print(json.dumps(topology, indent=2, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main())
